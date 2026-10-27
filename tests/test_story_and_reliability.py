"""Audit fixes (2026-10-21): loan numbers compared alike everywhere, one
delinquency rule, the Root Cause workbook complete and without duplicate
sheets, and the Dashboard's "why" and "recent loans" counting every loan."""
import inspect
import io

import pandas as pd
from openpyxl import load_workbook

from utils import delinquent_mask, loan_key, repair_upload, to_num


class TestLoanNumbers:
    def test_trimmed_upper_case_and_blank_kept(self):
        out = loan_key(pd.Series([" cvabc0123 ", "CVABC0124", None]))
        assert list(out[:2]) == ["CVABC0123", "CVABC0124"] and pd.isna(out[2])

    def test_whole_numbers_read_as_floats_lose_the_point_zero(self):
        # A loan number column with a blank is read as float: 12345.0 must still match "12345".
        assert list(loan_key(pd.Series([12345.0, None]))[:1]) == ["12345"]
        assert list(loan_key(pd.Series([12345, 678]))) == ["12345", "678"]

    def test_upload_cleans_them_and_says_how_many(self):
        # Real Sept closing: 16 lower-case loan numbers that never matched another file.
        df, notes = repair_upload(pd.DataFrame({"Loan No": ["cvabc0001", "CVABC0002", " CVABC0003"]}))
        assert list(df["Loan No"]) == ["CVABC0001", "CVABC0002", "CVABC0003"]
        assert any("Cleaned 2 loan number(s)" in n for n in notes)

    def test_numbers_turned_into_text_are_not_reported_as_faults(self):
        _, notes = repair_upload(pd.DataFrame({"Loan No": [101, 102]}))
        assert not any("loan number" in n for n in notes)

    def test_same_loan_differently_typed_is_one_loan_after_upload(self):
        from tests.test_utils import _build_upload
        from utils import load_and_validate
        # " l1" and "L1" are one loan: cleaned before duplicates are removed.
        upload = _build_upload({"Loan No": ["L1", " l1", "L2"]})
        df, errs = load_and_validate.__wrapped__(upload)
        assert errs == [] and sorted(df["Loan No"]) == ["L1", "L2"]
        assert df.attrs["dropped_duplicate_loans"] == 1


class TestOneDelinquencyRule:
    def test_mask_is_arrears_per_emi_above_zero_blank_not_behind(self):
        df = pd.DataFrame({"Arrears / EMI": [0, 0.01, None, -1, 3]})
        assert list(delinquent_mask(df)) == [False, True, False, False, True]

    def test_no_module_writes_the_rule_by_hand(self):
        import re
        import subprocess
        files = subprocess.run(["git", "ls-files", "*.py"], capture_output=True, text=True).stdout.split()
        hand = [f for f in files if not f.startswith("tests/") and f != "utils.py"
                and re.search(r'to_num\(\w+, "Arrears / EMI"\) > 0', open(f, encoding="utf-8-sig").read())]
        assert hand == []


def _master():
    df = pd.DataFrame({
        "Loan No": ["L1", "L2", "L3", "L4"], "RegionName": ["PUNE", "PUNE", "THANE", "THANE"],
        "Unit": ["BR1", "BR1", "BR2", None], "MNT NAME": ["A", "A", "B", "B"], "Zone": ["W", "W", "W", "W"],
        "Ag_Date": pd.to_datetime(["2026-01-01"] * 4), "Loan Status": ["RUN"] * 4,
        "Arrears / EMI": [0.5, 0, 7, 2], "Due Dt": [5, 5, 10, 10], "SOH": [1e6] * 4,
        "ARREARS AGAINST INST": [100, 0, 900, 300], "ARREARS AGAINST EXP": [0, 0, 0, 0],
    })
    return df


def _list():
    return pd.DataFrame({"MISMONTH": [20260916] * 2, "LOAN NO": ["L1", "L2"], "ARREARS / EMI": [1.5, 0.5]})


class TestRootCauseWorkbook:
    def _workbook(self):
        from analysis.root_cause import (build_root_cause_workbook, compute_recent_advances_bucket_movement,
                                         compute_recent_advances_daily_match, compute_recent_advances_story_by_grain)
        story = {g: compute_recent_advances_story_by_grain(_master(), _list(), grain=g) for g in ("region", "branch")}
        movement = {**story, "buckets": compute_recent_advances_bucket_movement(_master(), _list())}
        data = build_root_cause_workbook(pd.DataFrame({"Region": ["PUNE"]}), pd.DataFrame(), pd.DataFrame(), {},
                                         daily_match=compute_recent_advances_daily_match(_master(), _list()),
                                         movement=movement)
        return load_workbook(io.BytesIO(data))

    def test_one_sheet_per_level_no_duplicates(self):
        wb = self._workbook()
        assert "Recent Loans by Region" in wb.sheetnames and "Recent Loans by Branch" in wb.sheetnames
        grids = [tuple(tuple(c.value for c in r) for r in wb[n].iter_rows())
                 for n in wb.sheetnames if n.startswith(("Recent", "Bucket"))]
        assert len(grids) == len(set(grids)), "two sheets hold the same table"

    def test_sheets_use_the_words_on_screen(self):
        wb = self._workbook()
        heads = [c.value for c in wb["Recent Loans by Region"][1]]
        assert {"Loans", "Delinquent at month end", "Delinquent on list date", "Change (% points)",
                "Month-end cases cleared", "Newly delinquent"} <= set(heads)
        assert not {"Running Loans", "Close Delinquent", "Pending at Close", "New Slippage"} & set(heads)
        moves = [c.value for c in wb["Bucket Movement"][1]]
        assert moves[0] == "At month end" and "Not delinquent now" in moves

    def test_download_button_is_built_after_the_view(self):
        # Regression: built first, it handed out a workbook without the list figures.
        import ui.tabs.root_cause as rc
        src = inspect.getsource(rc.render_root_cause_tab)
        assert src.index("_render_recent_advances(") < src.index("st.download_button(")


class TestDashboardSnapshot:
    def test_why_counts_every_delinquent_loan_even_without_a_branch(self):
        from analysis.root_cause import ARREARS_GROUPS, root_cause_snapshot
        snap = root_cause_snapshot(_master())
        assert snap["delinquent"] == int(delinquent_mask(_master()).sum()) == 3      # L4 has no branch
        assert sum(snap["groups"][g][0] for g in ARREARS_GROUPS) == 3

    def test_recent_loans_behind_matches_month_end_arrears(self):
        from analysis.root_cause import root_cause_snapshot
        rec = root_cause_snapshot(_master())["recent"]
        assert rec["cohort_count"] == 4 and rec["behind"] == int((to_num(_master(), "Arrears / EMI") > 0).sum())
        # Same base as the Recent advances page: every running recent loan.
        assert rec["running"] == 4 and rec["behind_pct"] == 75.0

    def test_report_and_screen_share_one_label_per_group(self):
        from analysis.root_cause import GROUP_LABELS, WAS_SILENT
        from report_agent.story import SILENT_LABEL, _WHY_LABEL
        assert SILENT_LABEL == GROUP_LABELS[WAS_SILENT]
        assert all(_WHY_LABEL[k] == GROUP_LABELS[k] for k in _WHY_LABEL)


def test_tabs_follow_the_managers_questions():
    from ui.components import TAB_LABELS
    names = [t.split(" ", 1)[1] for t in TAB_LABELS]
    assert names == ["Dashboard", "Portfolio Intelligence", "Migration", "Root Cause", "Business",
                     "Action Lists", "AI Query", "Investigator", "Report"]


class TestNotPayingFromReceipts:
    """'Not paying 3 months' from Last Receipt Date (decision 2026-10-21): the
    file's own flag was Y for every 6+ EMI loan, including ones that paid."""

    def _book(self):
        receipts = (["2026-09-20"] * 60 + ["2026-08-10"] * 30 + ["2026-07-05", "2026-06-30", None, "2027-01-01"]
                    + ["2026-09-01"] * 6)
        n = len(receipts)
        return pd.DataFrame({
            "Loan No": [f"L{i}" for i in range(n)],
            "Last Receipt Date": pd.to_datetime(receipts),
            "Arrears / EMI": [1.0] * n,
            "No Coll 3 Months and >6 EMI": ["Y"] * n,
        })

    def test_window_is_the_three_months_up_to_the_files_latest_busy_month(self):
        from utils import receipt_window_start
        # One mistyped 2027 date must not move the window; Sept is the latest busy month.
        assert receipt_window_start(self._book()) == pd.Timestamp("2026-07-01")

    def test_behind_and_no_receipt_since_window_start(self):
        from config import NOT_PAYING_3M_FLAG_COL
        from utils import add_not_paying_flag
        df, note = add_not_paying_flag(self._book())
        flagged = df.loc[df[NOT_PAYING_3M_FLAG_COL] == "Y", "Last Receipt Date"]
        assert len(flagged) == 2                                   # 30 Jun, and never paid (blank)
        assert flagged.isna().sum() == 1 and (flagged.dropna() < "2026-07-01").all()
        assert "01 Jul 2026" in note and "not counted as not paying" in note

    def test_a_loan_not_behind_is_never_not_paying(self):
        from utils import not_paying_3m_mask
        book = self._book().assign(**{"Arrears / EMI": 0.0})
        assert not not_paying_3m_mask(book).any()

    def test_without_receipt_dates_the_files_flag_is_used_and_said(self):
        from config import NOT_PAYING_3M_FLAG_COL
        from utils import add_not_paying_flag
        df, note = add_not_paying_flag(self._book().drop(columns="Last Receipt Date"))
        assert (df[NOT_PAYING_3M_FLAG_COL] == "Y").all() and "file's own" in note

    def test_every_consumer_counts_the_same_loans(self):
        from analysis.root_cause import HARD_NOT_PAYING, WAS_SILENT, root_cause_snapshot
        from utils import add_not_paying_flag, loan_flags, not_paying_3m_mask
        df, _ = add_not_paying_flag(self._book().assign(**{"Arrears / EMI": [7.0] * 50 + [2.0] * 50}))
        n = int(not_paying_3m_mask(df).sum())
        assert int(loan_flags(df)["not_paying_3m"].sum()) == n
        groups = root_cause_snapshot(df.assign(Unit="BR1", RegionName="PUNE"))["groups"]
        assert groups[HARD_NOT_PAYING][0] + groups[WAS_SILENT][0] == n
