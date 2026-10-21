import io

import pandas as pd
import pytest
from openpyxl.utils import get_column_letter

from analysis.root_cause import (
    build_root_cause_workbook,
    clean_contaminated_flags,
    compute_recent_advances_bucket_summary,
    compute_recent_advances_daily_match,
    compute_recent_advances_status_by_grain,
    load_daily_missed_feed,
)


def _master():
    # Loan No, region, branch, executive (name, code), agreed, status
    rows = [
        ("L1", "PUNE", "BR1", "Amit Rao", "E1", "2025-12-01", "RUN"),
        ("L2", "PUNE", "BR1", "Amit Rao", "E1", "2026-01-01", "RUN"),
        ("L3", "PUNE", "BR2", "Sunil Jha", "E2", "2026-02-01", "RUN"),
        ("L4", "CS NAGAR", "BR3", "Amit Rao", "E3", "2025-11-15", "RUN"),   # same name, DIFFERENT branch
        ("L5", "CS NAGAR", "BR3", "Amit Rao", "E3", "2025-11-20", "RUN"),
        ("OLD", "PUNE", "BR1", "Amit Rao", "E1", "2025-06-01", "RUN"),      # before cohort start
        ("MAT1", "PUNE", "BR1", "Amit Rao", "E1", "2025-12-05", "MAT"),     # not running
    ]
    df = pd.DataFrame(rows, columns=["Loan No", "RegionName", "Unit", "MNT NAME", "MNT CODE", "Ag_Date", "Loan Status"])
    df["Ag_Date"] = pd.to_datetime(df["Ag_Date"])
    return df


def _feed():
    return pd.DataFrame({
        "LOAN NO": ["L1", "L2", "L4", "L5"],
        "REGIONNAME": ["PUNE", "PUNE", "CHHATRAPATI SAMBHAJI NAGAR", "CHHATRAPATI SAMBHAJI NAGAR"],
        "UNIT": ["BR1", "BR1", "BR3", "BR3"],
        "RE CODE": ["E1", "E1", "E3", "E3"],
        "RE NAME": ["Amit Rao "] * 4,
        "ARREARS / EMI": [0.5, 2.0, 3.0, 2.99],
    })


def _status(grain="region", master=None, feed=None):
    t = compute_recent_advances_status_by_grain(
        _master() if master is None else master, _feed() if feed is None else feed, grain=grain,
    )
    return t.set_index(t.columns[0])


class TestStatusTable:
    def test_running_loans_exclude_pre_cohort_and_non_running(self):
        t = _status()
        assert t.loc["PUNE", "Running Loans"] == 3          # L1, L2, L3 -- not OLD, not MAT1
        assert t.loc["CS NAGAR", "Running Loans"] == 2

    def test_bucket_split_uses_list_arrears_and_app_thresholds(self):
        t = _status()
        assert t.loc["PUNE", "Delinquent"] == 2
        assert t.loc["PUNE", "1-30 DPD"] == 1                # 0.5
        assert t.loc["PUNE", "SMA-2"] == 1                   # exactly 2.0
        assert t.loc["CS NAGAR", "NPA"] == 1                 # exactly 3.0
        assert t.loc["CS NAGAR", "SMA-2"] == 1               # 2.99

    def test_every_pct_is_of_running_loans(self):
        t = _status()
        assert t.loc["PUNE", "Delinquent %"] == 66.67
        assert t.loc["PUNE", "SMA-2 %"] == 33.33
        assert t.loc["Grand Total", "Delinquent %"] == 80.0  # 4 of 5

    def test_buckets_add_up_to_delinquent(self):
        t = _status()
        assert (t[["1-30 DPD", "SMA-1", "SMA-2", "NPA"]].sum(axis=1) == t["Delinquent"]).all()

    def test_list_loans_outside_the_running_cohort_are_never_counted(self):
        # OLD (pre-cohort), MAT1 (not running) and NEW (not in the LCC at all):
        # counting any of them would put loans in the numerator that aren't in
        # the denominator.
        feed = pd.concat([_feed(), pd.DataFrame({
            "LOAN NO": ["OLD", "MAT1", "NEW"], "REGIONNAME": ["PUNE"] * 3, "ARREARS / EMI": [1.0, 1.0, 1.0],
        })])
        t = _status(feed=feed)
        assert t.loc["PUNE", "Delinquent"] == 2
        assert t.loc["Grand Total", "Delinquent"] == 4
        assert (t["Delinquent"] <= t["Running Loans"]).all()

    def test_sidebar_filter_keeps_grand_total_consistent(self):
        # Regression: the list used to stay unfiltered, so filtering the LCC to
        # one region doubled the Grand Total Delinquent %.
        master = _master()
        t = _status(master=master[master["RegionName"] == "PUNE"])
        assert list(t.index) == ["PUNE", "Grand Total"]
        assert t.loc["Grand Total", "Delinquent %"] == t.loc["PUNE", "Delinquent %"]

    def test_region_comes_from_lcc_so_list_spelling_does_not_matter(self):
        t = _status()
        assert "CHHATRAPATI SAMBHAJI NAGAR" not in t.index
        assert t.loc["CS NAGAR", "Delinquent"] == 2

    def test_other_bucket_only_when_list_arrears_not_positive(self):
        assert "Other" not in _status().columns
        feed = _feed()
        feed.loc[0, "ARREARS / EMI"] = 0
        t = _status(feed=feed)
        assert t.loc["PUNE", "Other"] == 1
        assert t.loc["PUNE", "Delinquent"] == 2

    def test_duplicates_counted_once(self):
        master = pd.concat([_master(), _master()])
        t = _status(master=master, feed=pd.concat([_feed(), _feed()]))
        assert t.loc["PUNE", "Running Loans"] == 3
        assert t.loc["PUNE", "Delinquent"] == 2

    def test_branch_grain_carries_region(self):
        t = _status("branch")
        assert t.loc["BR1", "Region"] == "PUNE"
        assert t.loc["BR1", "Running Loans"] == 2
        assert t.loc["BR3", "NPA"] == 1

    def test_executive_grain_uses_composite_key_not_name_alone(self):
        t = compute_recent_advances_status_by_grain(_master(), _feed(), grain="executive")
        rows = t[t["Executive"] == "Amit Rao"].set_index("Branch")
        assert len(rows) == 2                                # same name, two branches
        assert rows.loc["BR1", "Running Loans"] == 2
        assert rows.loc["BR3", "NPA"] == 1

    def test_sorted_worst_first(self):
        t = compute_recent_advances_status_by_grain(_master(), _feed(), grain="branch")
        pct = t[t["Branch"] != "Grand Total"]["Delinquent %"].tolist()
        assert pct == sorted(pct, reverse=True)

    def test_list_missing_arrears_column_returns_empty(self):
        assert compute_recent_advances_status_by_grain(_master(), _feed().drop(columns=["ARREARS / EMI"])).empty

    def test_unknown_grain_raises(self):
        with pytest.raises(KeyError):
            compute_recent_advances_status_by_grain(_master(), _feed(), grain="planet")


class TestDailyMatchNotes:
    def _feed_with_strays(self):
        return pd.concat([_feed(), pd.DataFrame({
            "LOAN NO": ["OLD", "NEW", "FAR"],
            "REGIONNAME": ["PUNE", "PUNE", "THANE"],
            "UNIT": ["BR1", "BR9", "TH1"],
            "RE CODE": ["E1", "E9", "E8"],
            "ARREARS / EMI": [1.0, 1.0, 1.0],
        })], ignore_index=True)

    def test_every_list_loan_is_accounted_for(self):
        m = compute_recent_advances_daily_match(_master(), self._feed_with_strays())
        assert m["matched_count"] == 4
        assert m["in_view_not_counted_loan_nos"] == ["OLD"]
        assert m["missing_from_lcc_loan_nos"] == ["NEW"]
        assert m["missing_from_lcc_by_branch"] == {"BR9": 1}
        assert m["other_scope_count"] == 1                   # FAR: region not in this LCC
        parts = ("matched_count", "in_view_not_counted_count", "hidden_by_filter_count",
                 "missing_from_lcc_count", "other_scope_count")
        assert sum(m[p] for p in parts) == m["daily_feed_count"]

    def test_hidden_by_filter_is_separate_from_missing(self):
        master = _master()
        view = master[master["RegionName"] == "PUNE"]
        m = compute_recent_advances_daily_match(view, _feed(), df_all=master)
        assert m["matched_count"] == 2
        assert m["hidden_by_filter_count"] == 2              # L4, L5 are in CS NAGAR
        assert m["missing_from_lcc_count"] == 0

    def test_counts_executive_changes(self):
        feed = _feed()
        feed.loc[1, "RE CODE"] = "E7"
        assert compute_recent_advances_daily_match(_master(), feed)["exec_changed_count"] == 1

    def test_workbook_includes_status_and_notes(self):
        from openpyxl import load_workbook
        region = compute_recent_advances_status_by_grain(_master(), _feed(), grain="region")
        match = compute_recent_advances_daily_match(_master(), _feed())
        data = build_root_cause_workbook(
            pd.DataFrame({"Region": ["PUNE"]}), pd.DataFrame(), pd.DataFrame(), {},
            daily_match=match, region_status=region,
        )
        wb = load_workbook(io.BytesIO(data))
        assert "Recent Adv Region Status" in wb.sheetnames
        assert "Recent Adv Daily Feed Match" in wb.sheetnames


class TestLoadDailyFeed:
    def _xlsx(self, sheets: dict) -> io.BytesIO:
        buf = io.BytesIO()
        buf.name = "feed.xlsx"
        with pd.ExcelWriter(buf, engine="openpyxl") as w:
            for name, df in sheets.items():
                df.to_excel(w, sheet_name=name, index=False)
        buf.seek(0)
        return buf

    def test_skips_leading_pivot_sheet(self):
        pivot = pd.DataFrame({"Row Labels": ["PUNE"], "Count of LOAN NO": [2]})
        f = self._xlsx({"Sheet2": pivot, "Sheet1": _feed()})
        df, err = load_daily_missed_feed(f)
        assert err is None and len(df) == len(_feed()) and "LOAN NO" in df.columns

    def test_no_loan_column_returns_error(self):
        df, err = load_daily_missed_feed(self._xlsx({"Sheet1": pd.DataFrame({"a": [1]})}))
        assert df is None and "Loan No" in err


class TestRecentBucketsAndCleaning:
    def test_bucket_summary_uses_arrears_emi_thresholds(self):
        df = pd.DataFrame({
            "Loan No": list("abcdef"), "Ag_Date": ["2026-01-01"] * 6,
            "Arrears / EMI": [0, 0.5, 1.0, 2.0, 3.0, None], "SOH": [1e7] * 6,
        })
        b = compute_recent_advances_bucket_summary(df).set_index("Bucket")
        assert b["Count"].to_dict() == {"STD": 1, "0-1": 1, "1-2": 1, "2-3": 1, "NPA": 1}

    def test_contaminated_flag_values_become_null_and_are_counted(self):
        df = pd.DataFrame({"Non Starter": ["Yes", "No", "NEMCV", None], "CoLending_Loans": ["Y", "ALIVE", "N", "N"]})
        cleaned, contamination = clean_contaminated_flags(df)
        assert contamination == {"Non Starter": 1, "CoLending_Loans": 1}
        assert pd.isna(cleaned.loc[2, "Non Starter"]) and cleaned.loc[0, "Non Starter"] == "Yes"


# ── Month close vs list date ────────────────────────────────────────────────
from analysis.root_cause import (  # noqa: E402
    compute_recent_advances_bucket_movement,
    compute_recent_advances_then_vs_now_by_grain,
    compute_recent_advances_movement_by_grain,
    list_month_warning,
    list_report_date,
)


def _master_at_close():
    # L1, L2 pending at close; L3, L4, L5 clean. Due days 5/10/15/20/25.
    df = _master()
    df["Arrears / EMI"] = [0.5, 1.5, 0, 0, 0, 0, 0]
    df["Due Dt"] = [5, 10, 15, 20, 25, 5, 5]
    return df


def _list(day="20260912"):
    # L1 still pending (worse), L3 fresh (due 15th), L5 on the list though due the 25th.
    return pd.DataFrame({
        "MISMONTH": [day] * 3, "LOAN NO": ["L1", "L3", "L5"],
        "REGIONNAME": ["PUNE", "PUNE", "CS NAGAR"], "ARREARS / EMI": [1.2, 1.0, 0.3],
    })


def _move(grain="region", **kw):
    t = compute_recent_advances_movement_by_grain(_master_at_close(), _list(**kw), grain=grain)
    return t.set_index(t.columns[0])


class TestCloseVsList:
    def test_list_date_from_mismonth(self):
        assert list_report_date(_list()) == pd.Timestamp("2026-09-12")
        assert list_report_date(_list().assign(MISMONTH=20260912)) == pd.Timestamp("2026-09-12")
        assert list_report_date(_list().drop(columns="MISMONTH")) is None

    def test_pending_splits_into_still_pending_and_cleared(self):
        t = _move()
        assert t.loc["PUNE", "Pending at Close"] == 2
        assert t.loc["PUNE", "Still Pending"] == 1 and t.loc["PUNE", "Cleared"] == 1   # L1 on, L2 off
        assert t.loc["PUNE", "Cleared %"] == 50.0

    def test_new_slippage_only_counts_loans_already_due(self):
        # 12th: L3 (15th) and L4 (20th) aren't due yet, so PUNE has none due.
        t = _move()
        assert t.loc["PUNE", "Due by List Date"] == 0 and t.loc["PUNE", "New Slippage"] == 0
        # 16th: L3 is due and on the list -> new slippage; L4 (20th) still not due.
        t = _move(day="20260916")
        assert t.loc["PUNE", "Due by List Date"] == 1 and t.loc["PUNE", "New Slippage"] == 1
        assert t.loc["PUNE", "New Slippage %"] == 100.0

    def test_on_list_before_due_is_its_own_column(self):
        t = _move()
        assert t.loc["CS NAGAR", "On List, Not Yet Due"] == 1                     # L5, due the 25th
        assert "On List, Not Yet Due" not in _move(day="20260930").columns

    def test_every_loan_and_every_list_loan_counted_once(self):
        for day in ("20260912", "20260916", "20260930"):
            t = _move(day=day)
            early = t.get("On List, Not Yet Due", 0)
            assert (t["Pending at Close"] + t["Clean at Close"] == t["Running Loans"]).all()
            assert (t["Still Pending"] + t["Cleared"] == t["Pending at Close"]).all()
            status = _status(master=_master_at_close(), feed=_list(day=day))
            assert ((t["Still Pending"] + t["New Slippage"] + early).sort_index() == status["Delinquent"].sort_index()).all()

    def test_no_list_date_treats_every_loan_as_due(self):
        t = compute_recent_advances_movement_by_grain(_master_at_close(), _list().drop(columns="MISMONTH"))
        t = t.set_index("Region")
        assert t.loc["Grand Total", "Due by List Date"] == t.loc["Grand Total", "Clean at Close"]

    def test_month_warning(self):
        assert list_month_warning("2026-08", pd.Timestamp("2026-09-12")) is None
        assert "Sep 2026" in list_month_warning("2026-07", pd.Timestamp("2026-09-12"))
        assert "no MISMONTH" in list_month_warning("2026-08", None)

    def test_bucket_movement_rows_add_up_and_compare(self):
        m = compute_recent_advances_bucket_movement(_master_at_close(), _list()).set_index("Bucket at Close")
        assert (m[["Better", "Same", "Worse"]].sum(axis=1) == m["Total"]).all()
        assert m.loc["1-30 DPD", "SMA-1"] == 1 and m.loc["1-30 DPD", "Worse"] == 1    # L1: 0.5 -> 1.2
        assert m.loc["SMA-1", "Not on List"] == 1 and m.loc["SMA-1", "Better"] == 1   # L2 cleared
        assert m.loc["STD", "Worse"] == 2                                             # L3, L5
        assert m.loc["Total", "Total"] == 5

    def test_new_loan_with_no_due_day_is_its_own_group(self):
        # Real LCCs leave Due Dt and Arrears/EMI blank on loans with no EMI yet.
        # Counting them as due would pad the New Slippage denominator.
        master = _master_at_close()
        master.loc[master["Loan No"] == "L3", ["Due Dt", "Arrears / EMI"]] = [None, None]
        master.loc[master["Loan No"] == "L4", ["Due Dt", "Arrears / EMI"]] = [None, None]
        t = compute_recent_advances_movement_by_grain(master, _list(day="20260930")).set_index("Region")
        assert t.loc["PUNE", "New, No EMI at Close"] == 1 and t.loc["PUNE", "First EMI Missed"] == 1   # L3, on list
        assert t.loc["CS NAGAR", "New, No EMI at Close"] == 1 and t.loc["CS NAGAR", "First EMI Missed"] == 0  # L4
        assert t.loc["PUNE", "Clean at Close"] == 0 and t.loc["PUNE", "Due by List Date"] == 0
        assert (t["Pending at Close"] + t["Clean at Close"] + t["New, No EMI at Close"] == t["Running Loans"]).all()
        status = _status(master=master, feed=_list(day="20260930"))
        listed = t["Still Pending"] + t["New Slippage"] + t["First EMI Missed"]
        assert (listed.sort_index() == status["Delinquent"].sort_index()).all()
        m = compute_recent_advances_bucket_movement(master, _list(day="20260930")).set_index("Bucket at Close")
        assert m.loc["No EMI Yet", "Total"] == 2 and m.loc["No EMI Yet", "Worse"] == 1

    def test_new_loan_columns_hidden_when_none(self):
        assert "New, No EMI at Close" not in _move().columns

    def test_then_vs_now_figures_side_by_side(self):
        t = compute_recent_advances_then_vs_now_by_grain(_master_at_close(), _list()).set_index("Region")
        # PUNE: L1 (0.5 -> 1.2), L2 (1.5 -> cleared), L3 (clean -> 1.0). 3 running loans.
        assert t.loc["PUNE", "Close Delinquent"] == 2 and t.loc["PUNE", "Close 0-1"] == 1 and t.loc["PUNE", "Close 1-2"] == 1
        assert t.loc["PUNE", "List Delinquent"] == 2 and t.loc["PUNE", "List 1-2"] == 2
        assert t.loc["PUNE", "Close Delinquent %"] == 66.67 and t.loc["PUNE", "Change in Delinquent"] == 0
        # CS NAGAR: none at close, L5 on the list (0.3).
        assert t.loc["CS NAGAR", "Change in Delinquent % (pts)"] == 50.0
        assert t.loc["Grand Total", "List Delinquent"] == 3 and t.loc["Grand Total", "Running Loans"] == 5

    def test_then_vs_now_buckets_add_up_and_match_other_tables(self):
        master, lst = _master_at_close(), _list(day="20260930")
        t = compute_recent_advances_then_vs_now_by_grain(master, lst, grain="branch").set_index(["Branch", "Region"])
        bands = ["0-1", "1-2", "2-3", "NPA"]
        assert (t[[f"Close {b}" for b in bands]].sum(axis=1) == t["Close Delinquent"]).all()
        assert (t[[f"List {b}" for b in bands]].sum(axis=1) == t["List Delinquent"]).all()
        status = compute_recent_advances_status_by_grain(master, lst, grain="branch").set_index(["Branch", "Region"])
        move = compute_recent_advances_movement_by_grain(master, lst, grain="branch").set_index(["Branch", "Region"])
        assert (t["List Delinquent"].sort_index() == status["Delinquent"].sort_index()).all()
        assert (t["Close Delinquent"].sort_index() == move["Pending at Close"].sort_index()).all()

    def test_then_vs_now_has_no_verdict_column(self):
        t = compute_recent_advances_then_vs_now_by_grain(_master_at_close(), _list())
        assert not any("Status" in c or "Verdict" in c for c in t.columns)
        assert "List Other" not in t.columns                                          # none in this list

    def test_workbook_has_movement_sheets(self):
        from openpyxl import load_workbook
        movement = {"region": _move().reset_index(),
                    "buckets": compute_recent_advances_bucket_movement(_master_at_close(), _list())}
        data = build_root_cause_workbook(pd.DataFrame({"Region": ["PUNE"]}), pd.DataFrame(), pd.DataFrame(), {},
                                         movement=movement)
        names = load_workbook(io.BytesIO(data)).sheetnames
        assert "Close vs List Region" in names and "Bucket Movement" in names

    def test_shading_covers_every_unit_but_not_grand_total(self):
        # Regression: the range stopped one row early, leaving the last unit unshaded.
        from openpyxl import load_workbook
        region = compute_recent_advances_status_by_grain(_master(), _feed(), grain="region")   # 2 regions + total
        data = build_root_cause_workbook(pd.DataFrame({"Region": ["PUNE"]}), pd.DataFrame(), pd.DataFrame(), {},
                                         region_status=region)
        ws = load_workbook(io.BytesIO(data))["Recent Adv Region Status"]
        ranges = {str(r.sqref) for r in ws.conditional_formatting}
        col = get_column_letter(list(region.columns).index("Delinquent %") + 1)
        assert f"{col}2:{col}3" in ranges                  # rows 2-3 = both regions; row 4 = Grand Total
