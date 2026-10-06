"""Every Excel download goes through ui.components._excel_bytes, so this is the
one place the downloaded file's look is decided."""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from ui import components
from ui.components import _excel_bytes, style_excel_sheet


def _sheet(df):
    return load_workbook(BytesIO(_excel_bytes.__wrapped__(df))).active


def _scales(ws):
    """{cell range: (colour at min, colour at max)} for every colour-scale rule."""
    out = {}
    for rng, rules in ws.conditional_formatting._cf_rules.items():
        for rule in rules:
            colors = [c.rgb[-6:] for c in rule.colorScale.color]
            out[str(rng.sqref)] = (colors[0], colors[-1])
    return out


def _branch_table():
    return pd.DataFrame({
        "Branch": ["A", "B", "C", "Total"],
        "Accounts": [120, 80, 40, 240],
        "Collection %": [95.5, 88.0, 101.2, 94.1],
        "NPA %": [4.2, 9.8, 1.1, 5.0],
        "Loan Amount": [1250000.5, 830000.0, 410000.25, 2490000.75],
        "Year Of Manufacture": [2019, 2021, 2018, 0],
        "Share %": [50.0, 33.3, 16.7, 100.0],
    })


class TestHeaderAndLayout:
    def test_header_colours_freeze_and_filter(self):
        ws = _sheet(_branch_table())
        assert ws["A1"].fill.fgColor.rgb.endswith("1F4E78")              # blue header
        assert ws["A1"].font.color.rgb.endswith("FFFFFF") and ws["A1"].font.bold   # white text
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == "A1:G5"

    def test_total_row_is_bold_with_rule(self):
        ws = _sheet(_branch_table())
        assert ws["A5"].font.bold and ws["C5"].border.top.style == "medium"
        assert not ws["A4"].font.bold

    def test_columns_are_widened_to_fit(self):
        ws = _sheet(_branch_table())
        assert ws.column_dimensions["E"].width >= len("Loan Amount")


class TestColourScales:
    def test_direction_follows_the_metric_and_skips_total_row(self):
        scales = _scales(_sheet(_branch_table()))
        green, red, white = "63BE7B", "F8696B", "FFFFFF"
        assert scales["D2:D4"] == (white, red)   # NPA %: higher is worse, never green
        assert scales["C2:C4"] == (red, green)   # Collection %: higher is better

    def test_neutral_columns_get_no_colour(self):
        scales = _scales(_sheet(_branch_table()))
        assert not any(k.startswith(("B", "E", "G")) for k in scales)  # Accounts, Loan Amount, Share %


class TestNumberFormats:
    def test_formats(self):
        ws = _sheet(_branch_table())
        assert ws["C2"].number_format == '0.0"%"'            # percentages
        assert ws["E2"].number_format == "#,##0.00"          # amounts
        assert ws["F2"].number_format == "General"           # a year must not become 2,019
        assert ws["B2"].number_format == "General"           # small counts read fine

    def test_id_like_and_text_columns_untouched(self):
        ws = _sheet(pd.DataFrame({"Loan No": [100001, 100002], "Cust Name": ["a", "b"]}))
        assert ws["A2"].number_format == "General"

    def test_large_sheets_skip_per_cell_formats_but_keep_colours(self, monkeypatch):
        monkeypatch.setattr(components, "_XL_FORMAT_MAX_CELLS", 5)
        ws = _sheet(_branch_table())
        assert ws["C2"].number_format == "General"
        assert "D2:D4" in _scales(ws)

    def test_empty_frame_does_not_crash(self):
        _sheet(pd.DataFrame())
        _sheet(pd.DataFrame({"NPA %": []}))


def test_style_function_is_reusable_on_any_writer_sheet():
    from openpyxl import Workbook
    wb = Workbook()
    df = _branch_table()
    ws = wb.active
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    style_excel_sheet(ws, df)
    assert ws.freeze_panes == "A2"
